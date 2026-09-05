"""Create Mampffred's compact runtime catalog from the official BLS 4.0 XLSX.

The generated JSON is committed so normal builds never download or execute
upstream content. Run this script only after independently verifying the ZIP
digest documented in lib/bls-catalog.ts.
"""

from __future__ import annotations

import json
import hashlib
import math
import sys
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook


COLUMNS = {
    "energyKcal": "ENERCC Energie (Kilokalorien) [kcal/100g]",
    "proteinG": "PROT625 Protein (Nx6,25) [g/100g]",
    "fatG": "FAT Fett [g/100g]",
    "carbohydratesG": "CHO Kohlenhydrate, verfügbar [g/100g]",
    "fiberG": "FIBT Ballaststoffe, gesamt [g/100g]",
    "sugarG": "SUGAR Zucker (Mono- und Disaccharide), gesamt [g/100g]",
    "saltG": "NACL Salz (Natriumchlorid) [g/100g]",
}
SOURCE_SHA256 = "12b7a6ba62807ec9b301eb276f897dc85f99b2292311618dec3749a12d984c91"


def number_or_none(value: object, maximum: float) -> float | int | None:
    if not isinstance(value, (int, float)) or isinstance(value, bool):
        return None
    result = float(value)
    if not math.isfinite(result) or result < 0 or result > maximum:
        raise ValueError(f"invalid BLS nutrient value: {value!r}")
    return int(result) if result.is_integer() else result


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit("usage: extract-bls.py BLS_4_0_2025_DE.zip OUTPUT.json")
    source, destination = map(Path, sys.argv[1:])
    source_bytes = source.read_bytes()
    digest = hashlib.sha256(source_bytes).hexdigest()
    if digest != SOURCE_SHA256:
        raise ValueError(f"unexpected BLS source digest: {digest}")
    with ZipFile(BytesIO(source_bytes)) as archive:
        workbook_names = [
            name for name in archive.namelist() if name.endswith("BLS_4_0_Daten_2025_DE.xlsx")
        ]
        if len(workbook_names) != 1:
            raise ValueError("official BLS data workbook not found exactly once")
        workbook_bytes = archive.read(workbook_names[0])
    workbook = load_workbook(BytesIO(workbook_bytes), read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    header = next(rows)
    indexes = {str(value): index for index, value in enumerate(header)}
    required = ["BLS Code", "Lebensmittelbezeichnung", *COLUMNS.values()]
    missing = [column for column in required if column not in indexes]
    if missing:
        raise ValueError(f"missing BLS columns: {missing}")

    foods: list[list[object]] = []
    codes: set[str] = set()
    for row in rows:
        code = row[indexes["BLS Code"]]
        name = row[indexes["Lebensmittelbezeichnung"]]
        if not isinstance(code, str) or not isinstance(name, str):
            raise ValueError("BLS food without a code or German name")
        if code in codes:
            raise ValueError(f"duplicate BLS code: {code}")
        codes.add(code)
        values = [
            number_or_none(
                row[indexes[column]], 1_000 if key == "energyKcal" else 100
            )
            for key, column in COLUMNS.items()
        ]
        foods.append([code, name, *values])
    if len(foods) != 7_140:
        raise ValueError(f"expected 7140 BLS foods, got {len(foods)}")

    payload = {"schemaVersion": 1, "foods": foods}
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(
        json.dumps(payload, ensure_ascii=False, separators=(",", ":")) + "\n",
        encoding="utf-8",
    )


if __name__ == "__main__":
    main()
